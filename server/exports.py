"""Exports of what the user already sees: lists as CSV or Excel, documents as PDF or Word.

Spreadsheet cells never start a formula: text beginning with =, +, -, @ or a control character
gets a leading apostrophe (CSV) or is stored as plain text (Excel), so opening an export cannot
run anything a client, opponent or mailbox sender typed into a field.
"""
import csv
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Literal

from docx import Document
from docx.shared import Mm, Pt
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

from pathlib import Path

from .dates import display_date

# Full CI logo (artwork, name and claim) for the head of our own reports; drafts on the firm's letter paper never get it.
LOGO = Path(__file__).parent / 'assets' / 'schimmelpilz-logo.jpg'
LOGO_RATIO = 769 / 1984
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
FORMULA_START = ('=', '+', '-', '@', '\t', '\r', '\n')
ISO_DATE = re.compile(r'^\d{4}-\d{2}-\d{2}$')

STATUS = {'de': {'active': 'In Bearbeitung', 'waiting': 'Wartet', 'closed': 'Abgeschlossen'},
          'en': {'active': 'In progress', 'waiting': 'Waiting', 'closed': 'Closed'}}
AREA = {'de': {'general': 'Allgemein', 'consumer': 'Verbraucherrecht', 'tenancy': 'Mietrecht', 'contract': 'Vertragsrecht', 'commercial': 'Handelsrecht',
               'employment': 'Arbeitsrecht', 'family': 'Familienrecht', 'social': 'Sozialrecht', 'administrative': 'Verwaltungsrecht', 'criminal': 'Strafrecht'},
        'en': {'general': 'General', 'consumer': 'Consumer law', 'tenancy': 'Tenancy law', 'contract': 'Contract law', 'commercial': 'Commercial law',
               'employment': 'Employment law', 'family': 'Family law', 'social': 'Social law', 'administrative': 'Administrative law', 'criminal': 'Criminal law'}}
KIND = {'de': {'client': 'Mandant', 'opponent': 'Gegner', 'law_firm': 'Anwaltskanzlei', 'court': 'Gericht', 'authority': 'Behörde', 'company': 'Unternehmen', 'person': 'Privatperson', 'other': 'Sonstige'},
        'en': {'client': 'Client', 'opponent': 'Opponent', 'law_firm': 'Law firm', 'court': 'Court', 'authority': 'Authority', 'company': 'Company', 'person': 'Individual', 'other': 'Other'}}
ROLE = {'de': {'client': 'Mandant / eigene Partei', 'opponent': 'Gegenseite', 'counsel': 'Rechtsvertretung', 'court': 'Gericht', 'authority': 'Behörde', 'other': 'Weitere Beteiligte'},
        'en': {'client': 'Client / own party', 'opponent': 'Opposing party', 'counsel': 'Legal representative', 'court': 'Court', 'authority': 'Authority', 'other': 'Other party'}}


def day(value):
    """Stored ISO date or timestamp -> date for spreadsheet cells; None stays empty."""
    if not value:
        return None
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(value) if ISO_DATE.match(value) else datetime.fromisoformat(value).date()
    except ValueError:
        return value


def amount(value):
    try:
        return Decimal(value) if value not in (None, '') else None
    except (InvalidOperation, TypeError):
        return value


def safe_text(value):
    text = ILLEGAL_CHARACTERS_RE.sub('', str(value))
    return "'" + text if text.startswith(FORMULA_START) else text


def csv_cell(value, de):
    if value is None:
        return ''
    if isinstance(value, bool):
        return ('ja' if value else 'nein') if de else ('yes' if value else 'no')
    if isinstance(value, date):
        return display_date(value.isoformat(), de)
    if isinstance(value, (int, float, Decimal)):
        text = format(value, 'f') if isinstance(value, Decimal) else str(value)
        return text.replace('.', ',') if de else text
    return safe_text(value)


def csv_bytes(headers, rows, de):
    # Semicolons and a byte-order mark let German Excel open the file directly with umlauts intact.
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';' if de else ',', lineterminator='\r\n')
    writer.writerow(headers)
    writer.writerows([csv_cell(value, de) for value in row] for row in rows)
    return ('﻿' + output.getvalue()).encode('utf-8')


def xlsx_bytes(sheet, headers, rows, de):
    book = Workbook()
    table = book.active
    table.title = sheet[:31]
    table.append(headers)
    for cell in table[1]:
        cell.font = Font(bold=True)
    widths = [len(header) for header in headers]
    for index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            cell = table.cell(row=index, column=column)
            if value is None:
                continue
            if isinstance(value, bool):
                cell.value = ('ja' if value else 'nein') if de else ('yes' if value else 'no')
            elif isinstance(value, date):
                cell.value = value
                cell.number_format = 'DD.MM.YYYY' if de else 'DD/MM/YYYY'
            elif isinstance(value, (int, float, Decimal)):
                cell.value = value
                if isinstance(value, Decimal):
                    cell.number_format = '#,##0.00'
            else:
                cell.value = ILLEGAL_CHARACTERS_RE.sub('', str(value))
                cell.data_type = 's'  # never a formula, whatever the text starts with
            widths[column - 1] = max(widths[column - 1], min(60, max((len(part) for part in str(cell.value).split('\n')), default=0)))
    for column, width in enumerate(widths, start=1):
        table.column_dimensions[get_column_letter(column)].width = max(10, width + 2)
    table.freeze_panes = 'A2'
    if rows:
        table.auto_filter.ref = table.dimensions
    book.properties.creator = 'Schimmelpilz'
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def table_response(fmt, name, sheet, headers, rows, de):
    if fmt == 'csv':
        return Response(csv_bytes(headers, rows, de), media_type='text/csv; charset=utf-8', headers={'Content-Disposition': f'attachment; filename="{name}.csv"'})
    return Response(xlsx_bytes(sheet, headers, rows, de), media_type=XLSX, headers={'Content-Disposition': f'attachment; filename="{name}.xlsx"'})


def blocks_pdf(title, blocks):
    """blocks: (style, text) with style one of Title, Heading2, Heading3, BodyText."""
    styles, output = getSampleStyleSheet(), io.BytesIO()
    head = Table([[Image(str(LOGO), 62 * mm, 62 * mm * LOGO_RATIO)]], colWidths=[None], hAlign='LEFT')
    head.setStyle(TableStyle([('ALIGN', (0, 0), (-1, -1), 'LEFT'), ('LEFTPADDING', (0, 0), (-1, -1), 0), ('LINEBELOW', (0, 0), (-1, 0), 0.6, '#d9d2c4'), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
    story = [head, Spacer(1, 10)]
    for style, text in blocks:
        story.append(Paragraph(escape(str(text)).replace('\n', '<br/>'), styles[style]))
        story.append(Spacer(1, 6))
    SimpleDocTemplate(output, title=title).build(story)
    return output.getvalue()


def blocks_docx(title, blocks, note):
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    header = doc.sections[0].header.paragraphs[0]
    header.add_run().add_picture(str(LOGO), width=Mm(55))
    levels = {'Title': 0, 'Heading2': 2, 'Heading3': 3}
    for style, text in blocks:
        if style in levels:
            doc.add_heading(str(text), levels[style])
        else:
            lines = str(text).split('\n')
            paragraph = doc.add_paragraph(lines[0])
            for line in lines[1:]:
                paragraph.add_run().add_break()
                paragraph.add_run(line)
    doc.core_properties.title = title
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = note
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def document_response(fmt, name, title, blocks, note='Editable export. Changes made in Word are not reviewed by Schimmelpilz.'):
    if fmt == 'pdf':
        return Response(blocks_pdf(title, blocks), media_type='application/pdf', headers={'Content-Disposition': f'attachment; filename="{name}.pdf"'})
    return Response(blocks_docx(title, blocks, note), media_type=DOCX, headers={'Content-Disposition': f'attachment; filename="{name}.docx"'})


def case_rows(connection, user, lang):
    de = lang == 'de'
    headers = (['Aktenzeichen', 'Bezeichnung', 'Mandant/in', 'Gegenseite', 'Rechtsgebiet', 'Status', 'Zuständig', 'Gericht', 'Gerichtsaktenzeichen', 'Frist', 'Fristnotiz', 'Gegenstandswert', 'Währung', 'Dokumente', 'Entwürfe', 'Offene Aufgaben', 'Angelegt', 'Archiviert']
               if de else ['Reference', 'Title', 'Client', 'Opponent', 'Legal area', 'Status', 'Responsible', 'Court', 'Court reference', 'Deadline', 'Deadline note', 'Dispute value', 'Currency', 'Documents', 'Drafts', 'Open tasks', 'Created', 'Archived'])
    rows = []
    for case in visible_cases(connection, user):
        details = json.loads(case['details'])
        count = lambda sql: connection.execute(sql, (case['id'],)).fetchone()[0]
        rows.append([case['reference'], case['title'], details.get('client_name', ''), details.get('opponent_name', ''), AREA[lang].get(details.get('legal_area'), details.get('legal_area')),
                     STATUS[lang].get(details.get('status'), details.get('status')), details.get('responsible', ''), details.get('court', ''), details.get('court_reference', ''),
                     day(details.get('deadline')), details.get('deadline_note', ''), amount(f"{details['dispute_value']:.2f}") if details.get('dispute_value') is not None else None, details.get('currency', ''),
                     count('SELECT COUNT(*) FROM documents WHERE case_id=?'), count('SELECT COUNT(*) FROM drafts WHERE case_id=?'),
                     count('SELECT COUNT(*) FROM tasks WHERE case_id=? AND completed=0'), day(case['created_at']), day(case['archived_at']) if 'archived_at' in case.keys() else None])
    return headers, rows


def task_rows(connection, user, lang):
    de = lang == 'de'
    headers = (['Fällig am', 'Art', 'Aufgabe / Frist', 'Aktenzeichen', 'Akte', 'Priorität', 'Erledigt', 'Erledigt am', 'Notiz', 'Angelegt']
               if de else ['Due on', 'Type', 'Task / deadline', 'Reference', 'Case', 'Priority', 'Done', 'Done on', 'Note', 'Created'])
    priority = {'high': 'Hoch', 'normal': 'Normal', 'low': 'Niedrig'} if de else {'high': 'High', 'normal': 'Normal', 'low': 'Low'}
    rows = []
    for case in visible_cases(connection, user):
        details = json.loads(case['details'])
        if details.get('deadline'):
            rows.append([day(details['deadline']), 'Frist (Stammdaten)' if de else 'Deadline (case particulars)', details.get('deadline_note') or ('Frist' if de else 'Deadline'),
                         case['reference'], case['title'], None, None, None, '', None])
        for task in connection.execute('SELECT * FROM tasks WHERE case_id=? ORDER BY due_on IS NULL, due_on, created_at', (case['id'],)):
            rows.append([day(task['due_on']), 'Aufgabe' if de else 'Task', task['title'], case['reference'], case['title'], priority.get(task['priority'], task['priority']),
                         bool(task['completed']), day(task['completed_at']), task['note'], day(task['created_at'])])
    rows.sort(key=lambda row: (row[0] is None, str(row[0] or '')))
    return headers, rows


def contact_rows(connection, user, lang):
    de = lang == 'de'
    headers = (['Name', 'Art', 'z. Hd.', 'Anschrift', 'E-Mail', 'Telefon', 'Notiz', 'Akten', 'Geändert'] if de
               else ['Name', 'Type', 'Attention', 'Address', 'Email', 'Phone', 'Note', 'Cases', 'Updated'])
    rows = []
    for contact in connection.execute('SELECT * FROM contacts WHERE user_id=? ORDER BY name COLLATE NOCASE', (user['id'],)):
        links = connection.execute('SELECT c.reference, cc.role FROM contact_cases cc JOIN cases c ON c.id=cc.case_id WHERE cc.contact_id=? AND c.user_id=? ORDER BY c.reference', (contact['id'], user['id']))
        rows.append([contact['name'], KIND[lang].get(contact['kind'], contact['kind']), contact['attention'], contact['address'], contact['email'], contact['phone'], contact['note'],
                     ', '.join(f"{link['reference']} ({ROLE[lang].get(link['role'], link['role'])})" if link['role'] else link['reference'] for link in links), day(contact['updated_at'])])
    return headers, rows


def inbox_rows(connection, user, lang):
    de = lang == 'de'
    headers = (['Eingang am', 'Betreff', 'Absender', 'Datei', 'Größe (KB)', 'Aktenzeichen', 'Akte', 'Notiz', 'Hochgeladen'] if de
               else ['Received on', 'Subject', 'Sender', 'File', 'Size (KB)', 'Reference', 'Case', 'Note', 'Uploaded'])
    rows = []
    for item in connection.execute('''SELECT i.*, c.reference, c.title FROM inbox_documents i LEFT JOIN cases c ON c.id=i.case_id AND c.user_id=i.user_id
                                      WHERE i.user_id=? ORDER BY i.received_on DESC, i.created_at DESC''', (user['id'],)):
        rows.append([day(item['received_on']), item['subject'], item['sender'], item['original_name'], max(1, round(item['size_bytes'] / 1024)),
                     item['reference'] or ('ohne Zuordnung' if de else 'unassigned'), item['title'] or '', item['note'], day(item['created_at'])])
    return headers, rows


def visible_cases(connection, user):
    allowed = user.get('allowed_cases')
    return [case for case in connection.execute('SELECT * FROM cases WHERE user_id=? ORDER BY reference', (user['id'],)) if allowed is None or case['id'] in allowed]


DATASETS = {
    'cases': (case_rows, 'Schimmelpilz-Akten', {'de': 'Akten', 'en': 'Cases'}),
    'tasks': (task_rows, 'Schimmelpilz-Fristen', {'de': 'Fristen', 'en': 'Deadlines'}),
    'contacts': (contact_rows, 'Schimmelpilz-Kontakte', {'de': 'Kontakte', 'en': 'Contacts'}),
    'inbox': (inbox_rows, 'Schimmelpilz-Posteingang', {'de': 'Posteingang', 'en': 'Inbox'}),
}
# Case-scoped team members may only export lists built from the cases granted to them.
CASE_SCOPED = {'cases', 'tasks'}


def router(db, require_user, now):
    routes = APIRouter()

    @routes.get('/api/exports/{dataset}.{fmt}')
    def export(dataset: Literal['cases', 'tasks', 'contacts', 'inbox'], fmt: Literal['csv', 'xlsx'], lang: Literal['de', 'en'] = 'de', user: dict = Depends(require_user)):
        if user.get('allowed_cases') is not None and dataset not in CASE_SCOPED:
            raise HTTPException(403, 'case_scoped_member_use_case_workspace')
        build, name, sheet = DATASETS[dataset]
        with db() as connection:
            headers, rows = build(connection, user, lang)
        return table_response(fmt, f'{name}-{now()[:10]}', sheet[lang], headers, rows, lang == 'de')

    return routes
