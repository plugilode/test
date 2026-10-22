import csv
import io
import json
import uuid
from datetime import date

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from server import app as api
from server.tests import test_inbox
from server.tests.test_costs import invoice
from server.tests.test_route_plans import payload as route_payload


def read_csv(response):
    text = response.content.decode('utf-8')
    assert text.startswith('﻿')
    return list(csv.reader(io.StringIO(text[1:]), delimiter=';'))


def pdf_text(response):
    return ''.join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)


def docx_text(response):
    return '\n'.join(paragraph.text for paragraph in Document(io.BytesIO(response.content)).paragraphs)


class ExportTests(test_inbox.InboxTests):
    test_unassigned_persistence_then_exact_file_and_text_review_in_case = None
    test_ownership_validation_stale_updates_and_single_assignment = None
    test_rejected_uploads_leave_no_files_and_auth_guard_applies = None
    test_atomic_assignment_rollback_and_failed_upload_cleanup = None
    test_correction_guards_and_atomic_rollback = None
    test_correction_preserves_snapshots_blocks_old_sources_and_can_assign_back = None

    def test_lists_as_csv_and_excel_with_german_dates_and_no_formulas(self):
        with TestClient(api.app) as owner, TestClient(api.app) as other:
            self.login(owner); self.login(other, 'export-other@example.invalid')
            case = self.case(owner)
            details = {'client_name': '=HYPERLINK("http://evil.invalid","x")', 'legal_area': 'tenancy', 'status': 'waiting', 'deadline': '2026-11-05', 'deadline_note': 'Klagefrist'}
            current = owner.get(f'/api/cases/{case}').json()
            self.assertEqual(owner.put(f'/api/cases/{case}', json={'title': current['title'], 'description': current['description'], 'details': details}).status_code, 200)
            self.assertEqual(owner.post(f'/api/cases/{case}/tasks', json={'title': '+Stellungnahme prüfen', 'due_on': '2026-11-03', 'priority': 'high', 'note': ''}).status_code, 201)
            owner.post('/api/contacts', json={'name': '@Synthetic Kontakt', 'kind': 'client', 'case_links': [{'case_id': case, 'role': 'client'}]})
            self.upload(owner)
            other_case = self.case(other)

            rows = read_csv(owner.get('/api/exports/cases.csv'))
            self.assertEqual(rows[0][:3], ['Aktenzeichen', 'Bezeichnung', 'Mandant/in'])
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[1][2], '\'=HYPERLINK("http://evil.invalid","x")')
            self.assertEqual(rows[1][4:6], ['Mietrecht', 'Wartet'])
            self.assertEqual(rows[1][9], '05.11.2026')

            tasks = read_csv(owner.get('/api/exports/tasks.csv'))
            self.assertEqual([row[0] for row in tasks[1:]], ['03.11.2026', '05.11.2026'])
            self.assertEqual(tasks[1][2], "'+Stellungnahme prüfen")
            self.assertEqual(tasks[1][5], 'Hoch')
            english = owner.get('/api/exports/tasks.csv?lang=en').content.decode('utf-8')[1:]
            self.assertIn('03/11/2026', english)
            self.assertTrue(english.startswith('Due on,'))

            self.assertIn("'@Synthetic Kontakt", read_csv(owner.get('/api/exports/contacts.csv'))[1][0])
            inbox = read_csv(owner.get('/api/exports/inbox.csv'))
            self.assertEqual(inbox[1][:3], ['07.10.2026', 'Synthetic incoming invoice', 'Test Sender'])

            excel = owner.get('/api/exports/cases.xlsx')
            self.assertEqual(excel.status_code, 200)
            self.assertIn('Schimmelpilz-Akten-', excel.headers['content-disposition'])
            sheet = load_workbook(io.BytesIO(excel.content)).active
            self.assertEqual(sheet.title, 'Akten')
            cell = sheet.cell(row=2, column=3)
            self.assertEqual(cell.data_type, 's')
            self.assertEqual(cell.value, details['client_name'])
            self.assertEqual(sheet.cell(row=2, column=10).value.date(), date(2026, 11, 5))
            self.assertEqual(sheet.cell(row=2, column=10).number_format, 'DD.MM.YYYY')
            self.assertEqual(sheet.freeze_panes, 'A2')

            # Each account only exports its own data.
            self.assertEqual([row[0] for row in read_csv(other.get('/api/exports/cases.csv'))[1:]], [other.get(f'/api/cases/{other_case}').json()['reference']])
            self.assertEqual(len(read_csv(other.get('/api/exports/inbox.csv'))), 1)
            self.assertEqual(owner.get('/api/exports/drafts.csv').status_code, 422)
            self.assertEqual(owner.get('/api/exports/cases.json').status_code, 422)

    def test_case_scoped_member_exports_only_granted_cases(self):
        with TestClient(api.app) as owner, TestClient(api.app) as member:
            self.login(owner); self.login(member, 'export-member@example.invalid')
            me = owner.get('/api/auth/me').json(); person = member.get('/api/auth/me').json()
            first = self.case(owner); self.case(owner)
            token = owner.post('/api/team/invites', json={'email': person['email'], 'role': 'editor', 'case_ids': [first]}).json()['token']
            self.assertEqual(member.post('/api/team/join', json={'token': token}).status_code, 200)
            self.assertEqual(member.post('/api/team/switch', json={'owner_id': me['id']}).status_code, 200)
            rows = read_csv(member.get('/api/exports/cases.csv'))
            self.assertEqual([row[0] for row in rows[1:]], [owner.get(f'/api/cases/{first}').json()['reference']])
            self.assertEqual(member.get('/api/exports/tasks.xlsx').status_code, 200)
            self.assertEqual(member.get('/api/exports/contacts.csv').status_code, 403)
            self.assertEqual(member.get('/api/exports/inbox.xlsx').status_code, 403)

    def test_documents_as_pdf_and_word(self):
        with TestClient(api.app) as owner, TestClient(api.app) as other:
            self.login(owner); self.login(other, 'export-doc-other@example.invalid')
            case = self.case(owner); reference = owner.get(f'/api/cases/{case}').json()['reference']

            review = owner.post('/api/cost-reviews', json={'case_id': case, **invoice()}).json()
            word = owner.get(f'/api/cost-reviews/{review["id"]}/report.docx')
            self.assertEqual(word.status_code, 200)
            self.assertIn('Synthetic invoice check', docx_text(word))
            self.assertIn('Keine rechtliche Freigabe', docx_text(word))
            self.assertIn('Synthetic invoice check', pdf_text(owner.get(f'/api/cost-reviews/{review["id"]}/report.pdf')))
            lines = load_workbook(io.BytesIO(owner.get(f'/api/cost-reviews/{review["id"]}/lines.xlsx').content)).active
            self.assertEqual(lines.cell(row=2, column=2).value, 'Test service')
            self.assertEqual(float(lines.cell(row=2, column=5).value), 200.0)
            self.assertIn('Errechnet zahlbar', read_csv(owner.get(f'/api/cost-reviews/{review["id"]}/lines.csv'))[-3][1])
            self.assertEqual(other.get(f'/api/cost-reviews/{review["id"]}/report.docx').status_code, 404)
            self.assertEqual(other.get(f'/api/cost-reviews/{review["id"]}/lines.csv').status_code, 404)

            plan = owner.put(f'/api/cases/{case}/route-plan', json=route_payload(expected_version=0))
            self.assertEqual(plan.status_code, 200, plan.text)
            self.assertIn(reference, docx_text(owner.get(f'/api/cases/{case}/route-plan/report.docx')))
            self.assertEqual(other.get(f'/api/cases/{case}/route-plan/report.docx').status_code, 404)

            run_id = str(uuid.uuid4())
            legal = [{'citation_id': 'L1', 'title': '§ 536 BGB · Mietminderung', 'pinpoint': 'Abs. 1', 'url': 'https://www.gesetze-im-internet.de/bgb/__536.html', 'fetched_at': '2026-10-21T01:00:00+00:00', 'text': 'Synthetic statute text.'}]
            with api.db() as connection:
                connection.execute('''INSERT INTO research_runs(id,case_id,question,language,workflow,answer,status,model,created_at,finished_at,sources,document_selection,legal_sources,legal_selection,citation_audit,retrieval_report)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (run_id, case, 'Synthetic question?', 'de', 'research', '**Ja.** Grund [L1].\n- Schritt eins', 'completed', 'test', '2026-10-22T08:00:00+00:00', '2026-10-22T08:00:05+00:00', '[]', '[]', json.dumps(legal), '{}', '{}', '{}'))
            answer = docx_text(owner.get(f'/api/cases/{case}/research/{run_id}/answer.docx'))
            self.assertIn('Synthetic question?', answer)
            self.assertIn('Ja. Grund [L1].', answer)
            self.assertIn('• Schritt eins', answer)
            self.assertIn('[L1] § 536 BGB · Mietminderung · Abs. 1', answer)
            self.assertIn('22.10.2026', answer)
            self.assertIn('Synthetic statute text.', pdf_text(owner.get(f'/api/cases/{case}/research/{run_id}/answer.pdf')))
            self.assertEqual(other.get(f'/api/cases/{case}/research/{run_id}/answer.pdf').status_code, 404)
            self.assertEqual(owner.get(f'/api/cases/{case}/research/{uuid.uuid4()}/answer.pdf').status_code, 404)
